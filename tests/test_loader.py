import pytest

from src.ingest.loader import load_document
from src.utils.exceptions import LoadError, UnsupportedFormatError


def test_txt_single_page(tmp_path, settings):
    p = tmp_path / "a.md"; p.write_text("# 标题\n\n正文", encoding="utf-8")
    doc = load_document(p, settings)
    assert doc.pages[0].page_no == 1 and "正文" in doc.pages[0].text

def test_unsupported_ext_detail(tmp_path, settings):
    p = tmp_path / "a.pdf"; p.write_bytes(b"%PDF-1.4")
    with pytest.raises(UnsupportedFormatError) as ei:
        load_document(p, settings)
    assert ei.value.detail["ext"] == ".pdf"
    assert ".md" in ei.value.detail["supported"]

def test_zero_byte_raises(tmp_path, settings):
    p = tmp_path / "a.txt"; p.write_bytes(b"")
    with pytest.raises(LoadError):
        load_document(p, settings)


def test_docx_paragraphs_are_read(tmp_path, settings):
    """docx 正文段落必须能读出来（曾经因 lxml 命名空间写法错误而全空）。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("差旅住宿标准")
    doc.add_paragraph("一线城市每晚不超过 600 元")
    path = tmp_path / "a.docx"
    doc.save(str(path))

    raw = load_document(path, settings)
    assert "差旅住宿标准" in raw.pages[0].text
    assert "600 元" in raw.pages[0].text